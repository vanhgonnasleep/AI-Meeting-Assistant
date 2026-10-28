import io
import docx
from database.models import MeetingRecord, ActionItem
from app.pipeline.exporter import generate_meeting_docx


def test_generate_meeting_docx_structure():
    record = MeetingRecord(
        filename="Q3_Executive_Strategy.mp3",
        raw_transcript="Alex: We are ready to launch. Sarah: Budget is confirmed.",
        executive_summary="- Agreed on Q3 rollout plan.\n- Budget locked at $50,000.",
        action_items=[
            ActionItem(task="Deploy production cluster", assignee="Alex", deadline="2026-11-15", status="pending"),
            ActionItem(task="Review marketing assets", assignee="Sarah", deadline="2026-11-10", status="completed"),
        ],
        duration=1845.5,
        language="vi",
        insights={
            "decisions": [{"text": "Chốt ngày phát hành 15/11", "speaker": "Alex", "timestamp": "12:30"}],
            "risks": [{"text": "Nguy cơ nghẽn tải máy chủ", "speaker": "Sarah", "timestamp": "15:45"}],
            "open_questions": []
        }
    )

    buf = generate_meeting_docx(record)
    assert isinstance(buf, io.BytesIO)
    assert buf.getbuffer().nbytes > 1000

    # Verify it can be loaded by python-docx without error
    doc = docx.Document(buf)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "BIÊN BẢN CUỘC HỌP" in text
    assert "Q3_Executive_Strategy.mp3" in " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "Chốt ngày phát hành 15/11" in text
    assert "NGƯỜI LẬP BIÊN BẢN" in " ".join(cell.text for row in doc.tables[-1].rows for cell in row.cells)
