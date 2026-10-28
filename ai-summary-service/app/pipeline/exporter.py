"""Corporate Meeting Minutes Export Engine (.docx generation).

Generates standardized enterprise-ready meeting minutes with styled metadata,
bulleted executive summary, structured action deliverables table, and signature blocks.
"""
import io
from datetime import datetime
from typing import Any, Dict, List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn


def _set_cell_background(cell, fill_hex: str) -> None:
    """Sets background color of a Word table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def _set_table_borders(table, color_hex: str = "D1D5DB") -> None:
    """Sets subtle, clean table borders."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
            f'  <w:insideV w:val="none"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)


def generate_meeting_docx(meeting: Any) -> io.BytesIO:
    """
    Compiles a meeting record into a high-polish Microsoft Word (.docx) document.
    """
    doc = docx.Document()

    # 1. Page Margins (Standard 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Document Title & Subtitle
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("BIÊN BẢN CUỘC HỌP / MEETING MINUTES")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)  # slate-900

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    sub_run = sub_p.add_run(f"Hệ thống trợ lý cuộc họp Clarion AI · Mã lưu trữ: #{getattr(meeting, 'id', 'DRAFT')}")
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)  # slate-500

    # 3. Overview Metadata Card (2-column table)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_fields = [
        ("Tên tài liệu / Title:", str(getattr(meeting, "filename", "Cuộc họp"))),
        ("Thời gian / Timestamp:", str(getattr(meeting, "created_at", datetime.now().strftime("%Y-%m-%d %H:%M")))),
        ("Thời lượng / Duration:", f"{round(float(getattr(meeting, 'duration', 0) or 0) / 60, 1)} phút ({round(float(getattr(meeting, 'duration', 0) or 0), 1)}s)"),
        ("Ngôn ngữ / Language:", "Tiếng Việt (vi)" if getattr(meeting, "language", "") == "vi" else "English (en)"),
    ]

    for idx, (label, val) in enumerate(meta_fields):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        _set_cell_background(cell_lbl, "F8FAFC")
        _set_cell_background(cell_val, "FFFFFF")

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(2)
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(71, 85, 105)

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(2)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = RGBColor(15, 23, 42)

    _set_table_borders(meta_table, "E2E8F0")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. Section I: Executive Summary
    h1 = doc.add_heading("I. TỔNG QUAN ĐIỀU HÀNH (EXECUTIVE SUMMARY)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    for run in h1.runs:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    summary_text = getattr(meeting, "executive_summary", "") or ""
    if summary_text.strip():
        for line in summary_text.split("\n"):
            clean_line = line.strip()
            if not clean_line:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            if clean_line.startswith(("-", "*", "•")):
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("•  " + clean_line.lstrip("-*• \t"))
            else:
                run = p.add_run(clean_line)
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
    else:
        empty_p = doc.add_paragraph("(Chưa có nội dung tóm tắt)")
        empty_p.runs[0].font.italic = True
        empty_p.runs[0].font.size = Pt(10)

    # 5. Section II: Key Decisions & Risks Matrix
    insights = getattr(meeting, "insights", {}) or {}
    decisions = insights.get("decisions", []) if isinstance(insights, dict) else []
    risks = insights.get("risks", []) if isinstance(insights, dict) else []

    if decisions or risks:
        h2 = doc.add_heading("II. QUYẾT ĐỊNH & RỦI RO (GOVERNANCE MATRIX)", level=1)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(6)
        for run in h2.runs:
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)

        if decisions:
            doc.add_paragraph("1. Quyết định trọng yếu (Key Decisions):").runs[0].font.bold = True
            for dec in decisions:
                txt = dec.get("text") if isinstance(dec, dict) else str(dec)
                spk = dec.get("speaker") if isinstance(dec, dict) else None
                ts = dec.get("timestamp") if isinstance(dec, dict) else None
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(f"✓  {txt}")
                r.font.size = Pt(10)
                if spk or ts:
                    extra = f" [{spk or ''}{' · ' + ts if ts else ''}]"
                    r_ex = p.add_run(extra)
                    r_ex.font.size = Pt(9)
                    r_ex.font.color.rgb = RGBColor(100, 116, 139)

        if risks:
            p_risk_h = doc.add_paragraph("2. Rủi ro & Điểm nghẽn (Risks & Blockers):")
            p_risk_h.paragraph_format.space_before = Pt(6)
            p_risk_h.runs[0].font.bold = True
            for risk in risks:
                txt = risk.get("text") if isinstance(risk, dict) else str(risk)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(f"⚠  {txt}")
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(180, 83, 9)

    # 6. Section III: Action Items Table
    action_items = getattr(meeting, "action_items", []) or []
    h3 = doc.add_heading("III. PHÂN CÔNG CÔNG VIỆC (ACTION ITEMS)", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    for run in h3.runs:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    if action_items:
        table = doc.add_table(rows=len(action_items) + 1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = ["STT", "Nội dung đầu việc (Deliverable)", "Phụ trách", "Hạn chót", "Trạng thái"]
        col_widths = [Inches(0.6), Inches(3.2), Inches(1.3), Inches(1.1), Inches(1.1)]

        header_row = table.rows[0]
        for col_idx, (text, width) in enumerate(zip(headers, col_widths)):
            cell = header_row.cells[col_idx]
            cell.width = width
            _set_cell_background(cell, "0F172A")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, item in enumerate(action_items, start=1):
            row = table.rows[row_idx]
            task_str = getattr(item, "task", "") if hasattr(item, "task") else (item.get("task", "") if isinstance(item, dict) else str(item))
            assignee_str = getattr(item, "assignee", "") if hasattr(item, "assignee") else (item.get("assignee", "Chưa giao") if isinstance(item, dict) else "Chưa giao")
            deadline_str = getattr(item, "deadline", "") if hasattr(item, "deadline") else (item.get("deadline", "Chưa có") if isinstance(item, dict) else "Chưa có")
            status_str = getattr(item, "status", "pending") if hasattr(item, "status") else (item.get("status", "pending") if isinstance(item, dict) else "pending")

            vals = [
                str(row_idx),
                task_str,
                assignee_str or "Chưa giao",
                deadline_str or "Chưa có",
                "Đã xong" if status_str.lower() in ("completed", "done") else "Đang làm"
            ]

            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"

            for col_idx, (val, width) in enumerate(zip(vals, col_widths)):
                cell = row.cells[col_idx]
                cell.width = width
                _set_cell_background(cell, bg_color)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (0, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.font.size = Pt(9)
                if col_idx == 4 and status_str.lower() in ("completed", "done"):
                    run.font.color.rgb = RGBColor(22, 101, 52)
                    run.font.bold = True
                else:
                    run.font.color.rgb = RGBColor(30, 41, 59)

        _set_table_borders(table, "CBD5E1")
    else:
        doc.add_paragraph("(Không có đầu việc cần thực hiện)")

    # 7. Section IV: Corporate Signature Block
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    h4 = doc.add_heading("IV. XÁC NHẬN & CHỮ KÝ (APPROVAL SIGNATURES)", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(8)
    for run in h4.runs:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.rows[0].cells[0].width = Inches(3.5)
    sig_table.rows[0].cells[1].width = Inches(3.5)

    p_sig1 = sig_table.rows[0].cells[0].paragraphs[0]
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_sig1.add_run("NGƯỜI LẬP BIÊN BẢN\n(Ký và ghi rõ họ tên)")
    r1.font.bold = True
    r1.font.size = Pt(10)

    p_sig2 = sig_table.rows[0].cells[1].paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p_sig2.add_run("CHỦ TỌA CUỘC HỌP\n(Ký duyệt xác nhận)")
    r2.font.bold = True
    r2.font.size = Pt(10)

    # Blank spacing for actual signature
    sig_table.rows[1].cells[0].paragraphs[0].add_run("\n\n\n\n_______________________").font.color.rgb = RGBColor(148, 163, 184)
    sig_table.rows[1].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_table.rows[1].cells[1].paragraphs[0].add_run("\n\n\n\n_______________________").font.color.rgb = RGBColor(148, 163, 184)
    sig_table.rows[1].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
